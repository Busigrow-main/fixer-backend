#!/usr/bin/env python3
"""Build a Google-Docs-friendly Word file from fixxer-shop-api.html."""

from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
HTML_PATH = HERE / "fixxer-shop-api.html"
DOCX_PATH = HERE / "Fixxer-Shop-API.docx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
STEEL = RGBColor(0x2E, 0x75, 0xB6)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1F4E79"
CODE_FILL = "F4F4F4"
ALT_ROW = "F7F9FC"
BORDER = "8AA0B8"

METHOD_COLORS = {
    "GET": RGBColor(0x0D, 0x73, 0x2E),
    "POST": RGBColor(0x1F, 0x4E, 0x79),
    "PUT": RGBColor(0xC0, 0x56, 0x00),
    "PATCH": RGBColor(0x7B, 0x2D, 0x8E),
    "DELETE": RGBColor(0xC0, 0x39, 0x2B),
}


def set_run_font(run, name="Calibri", size=11, bold=False, color=INK, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.08):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade_cell(cell, fill):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tcmar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcmar.append(node)
    tcpr.append(tcmar)


def set_table_borders(table, color=BORDER, size="8"):
    tbl = table._tbl
    tblpr = tbl.tblPr
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def set_table_width(table):
    tbl = table._tbl
    tblpr = tbl.tblPr
    for old in tblpr.findall(qn("w:tblW")):
        tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), "5000")
    tblw.set(qn("w:type"), "pct")
    tblpr.append(tblw)
    look = OxmlElement("w:tblLook")
    look.set(qn("w:firstRow"), "1")
    look.set(qn("w:noHBand"), "0")
    look.set(qn("w:noVBand"), "1")
    tblpr.append(look)


def prevent_row_split(row):
    tr = row._tr
    trpr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trpr.append(cant)


def style_heading(paragraph, level):
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    if level == 0:
        set_run_font(run, size=26, bold=True, color=NAVY)
        set_paragraph_spacing(paragraph, before=0, after=6, line=1.15)
    elif level == 1:
        set_run_font(run, size=16, bold=True, color=NAVY)
        set_paragraph_spacing(paragraph, before=18, after=8, line=1.15)
        paragraph.paragraph_format.keep_with_next = True
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=STEEL)
        set_paragraph_spacing(paragraph, before=14, after=6, line=1.15)
        paragraph.paragraph_format.keep_with_next = True
    else:
        set_run_font(run, size=11, bold=True, color=MUTED)
        set_paragraph_spacing(paragraph, before=10, after=4, line=1.1)
        paragraph.paragraph_format.keep_with_next = True


def add_bottom_border(paragraph, color="1F4E79", size="12"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def add_method_heading(doc, text):
    paragraph = doc.add_heading("", level=2)
    paragraph.clear()
    parts = text.strip().split(" ", 1)
    method = parts[0]
    path = parts[1] if len(parts) > 1 else ""
    if method in METHOD_COLORS:
        mrun = paragraph.add_run(method)
        set_run_font(mrun, size=13, bold=True, color=METHOD_COLORS[method])
        if path:
            prun = paragraph.add_run("  " + path)
            set_run_font(prun, size=13, bold=True, color=STEEL)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=13, bold=True, color=STEEL)
    set_paragraph_spacing(paragraph, before=14, after=6, line=1.15)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_rich_paragraph(doc, node):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=0, after=8, line=1.15)
    if not node:
        return paragraph

    def walk(el, bold=False):
        if isinstance(el, NavigableString):
            text = str(el)
            if text:
                run = paragraph.add_run(text)
                set_run_font(run, size=11, bold=bold, color=INK)
            return
        if not isinstance(el, Tag):
            return
        next_bold = bold or el.name in ("b", "strong")
        if el.name == "br":
            paragraph.add_run().add_break()
            return
        for child in el.children:
            walk(child, next_bold)

    walk(node, False)
    return paragraph


def fill_cell_text(cell, text, *, header=False, code=False, first_col_label=False):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=0, after=0, line=1.08)
    run = p.add_run(text.strip())
    if header:
        set_run_font(run, size=10, bold=True, color=WHITE)
    elif code:
        set_run_font(run, name="Courier New", size=9, color=INK)
    elif first_col_label:
        set_run_font(run, size=10, bold=True, color=NAVY)
    else:
        set_run_font(run, size=10, color=INK)
    set_cell_margins(cell)


def add_html_table(doc, table_el):
    rows = table_el.find_all("tr", recursive=False)
    if not rows:
        rows = table_el.find_all("tr")
    if not rows:
        return
    ncol = max(len(r.find_all(["th", "td"], recursive=False)) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    set_table_borders(table)
    set_table_width(table)

    header_row = bool(rows[0].find("th")) and all(
        c.name == "th" for c in rows[0].find_all(["th", "td"], recursive=False)
    )
    label_col = (not header_row) and any(r.find("th") for r in rows)

    for r_i, row_el in enumerate(rows):
        cells_el = row_el.find_all(["th", "td"], recursive=False)
        prevent_row_split(table.rows[r_i])
        for c_i in range(ncol):
            cell = table.cell(r_i, c_i)
            text = cells_el[c_i].get_text(" ", strip=True) if c_i < len(cells_el) else ""
            is_th = c_i < len(cells_el) and cells_el[c_i].name == "th"
            is_header = header_row and r_i == 0
            fill_cell_text(
                cell,
                text,
                header=is_header,
                first_col_label=(label_col and is_th and not is_header),
            )
            if is_header:
                shade_cell(cell, HEADER_FILL)
            elif label_col and is_th:
                shade_cell(cell, "E8EEF4")
            elif not header_row and not label_col and r_i % 2 == 1:
                shade_cell(cell, ALT_ROW)
            elif header_row and r_i > 0 and r_i % 2 == 0:
                shade_cell(cell, ALT_ROW)

    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=6, line=1.0)


def add_code_block(doc, text):
    text = (text or "").strip("\n")
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_borders(table, color="D0D5DD", size="4")
    set_table_width(table)
    cell = table.cell(0, 0)
    shade_cell(cell, CODE_FILL)
    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
    cell.text = ""
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.keep_together = True
        set_paragraph_spacing(p, before=0, after=0, line=1.05)
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Courier New", size=9, color=RGBColor(0x1A, 0x1A, 0x1A))
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=8, line=1.0)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing = 1.15

    for style_name, size, color, before, after in (
        ("Title", 26, NAVY, 0, 6),
        ("Heading 1", 16, NAVY, 18, 8),
        ("Heading 2", 13, STEEL, 14, 6),
        ("Heading 3", 11, MUTED, 10, 4),
    ):
        try:
            st = doc.styles[style_name]
        except KeyError:
            continue
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)


def add_footer(doc):
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Fixxer Shop API  ·  Page ")
    set_run_font(run, size=9, color=MUTED)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r2 = p.add_run()
    r2._r.append(fld1)
    r2._r.append(instr)
    r2._r.append(fld2)
    set_run_font(r2, size=9, color=MUTED)


def convert():
    soup = BeautifulSoup(HTML_PATH.read_text(encoding="utf-8"), "html.parser")
    body = soup.body
    doc = Document()
    configure_styles(doc)
    add_footer(doc)

    for el in body.children:
        if isinstance(el, NavigableString):
            continue
        if not isinstance(el, Tag):
            continue
        name = el.name
        if name == "h1":
            p = doc.add_heading(el.get_text(strip=True), level=0)
            style_heading(p, 0)
            add_bottom_border(p, "1F4E79", "18")
        elif name == "h2":
            p = doc.add_heading(el.get_text(strip=True), level=1)
            style_heading(p, 1)
            add_bottom_border(p, "8AA0B8", "8")
        elif name == "h3":
            text = el.get_text(" ", strip=True)
            if text.split(" ", 1)[0] in METHOD_COLORS:
                add_method_heading(doc, text)
            else:
                p = doc.add_heading(text, level=2)
                style_heading(p, 2)
        elif name == "h4":
            p = doc.add_heading(el.get_text(strip=True), level=3)
            style_heading(p, 3)
        elif name == "p":
            add_rich_paragraph(doc, el)
        elif name == "table":
            add_html_table(doc, el)
        elif name == "pre":
            add_code_block(doc, el.get_text())

    doc.save(DOCX_PATH)
    print(f"Wrote {DOCX_PATH} ({DOCX_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    convert()
